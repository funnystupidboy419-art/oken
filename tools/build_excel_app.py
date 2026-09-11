#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
通信表・指導要録 所見作成支援  Excel版ワークブックの生成

  python3 tools/build_excel_app.py  ->  dist/所見作成支援.xlsx

生成したブックに vba/*.bas を取り込んで .xlsm として保存すると、
マクロ付きアプリになります（手順は「はじめに」シートに記載）。
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

OUT_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "dist")
OUT_PATH = os.path.join(OUT_DIR, "所見作成支援.xlsx")

ACCENT = "2F6F4F"
HEAD_FILL = PatternFill("solid", fgColor="EAF3EE")
TITLE_FONT = Font(bold=True, size=14, color=ACCENT)
HEAD_FONT = Font(bold=True, size=11)
NOTE_FONT = Font(size=10, color="5B6B7C")
THIN = Side(style="thin", color="D7DEE8")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

MARK_A = "◎"
MARK_B = "○"

BEHAVIOR_ITEMS = [
    ("seikatsu", "基本的な生活習慣"),
    ("kenko", "健康・体力の向上"),
    ("jishu", "自主・自律"),
    ("sekinin", "責任感"),
    ("soui", "創意工夫"),
    ("omoiyari", "思いやり・協力"),
    ("seimei", "生命尊重・自然愛護"),
    ("kinro", "勤労・奉仕"),
    ("kosei", "公正・公平"),
    ("kokyo", "公共心・公徳心"),
]

CATEGORIES = [
    ("observation", "日常の見取り"),
    ("event", "行事・行事の係"),
    ("committee", "委員会"),
    ("kakari", "係"),
    ("behavior", "行動の記録"),
]

MORAL_VALUES = [
    "善悪の判断、自律、自由と責任", "正直、誠実", "節度、節制", "個性の伸長",
    "希望と勇気、努力と強い意志", "真理の探究",
    "親切、思いやり", "感謝", "礼儀", "友情、信頼", "相互理解、寛容",
    "規則の尊重", "公正、公平、社会正義", "勤労、公共の精神",
    "家族愛、家庭生活の充実", "よりよい学校生活、集団生活の充実",
    "伝統と文化の尊重、国や郷土を愛する態度", "国際理解、国際親善",
    "生命の尊さ", "自然愛護", "感動、畏敬の念", "よりよく生きる喜び",
]

OBS_PATTERNS = [
    ("sugata", "〜する姿が見られました", "{{見取り}}姿が{見られました|見られた}。",
     "例）友達の失敗を責めずに励ます"),
    ("dekita", "〜ことができました", "{{見取り}}ことができ{ました|た}。",
     "例）自分の考えを理由を付けて発表する"),
    ("teita", "〜ていました", "{{見取り}}{ていました|ていた}。",
     "例）当番の仕事を人一倍ていねいに行っ"),
    ("datta", "〜でした", "{{見取り}}{でした|であった}。",
     "例）学級のムードメーカー"),
    ("as-is", "そのまま使う（要録は自動で常体に）", "{{見取り}}",
     "例）友達の意見をよく聞き、話し合いをまとめていました。"),
]

# ---------------------------------------------------------------------------
# 文例  (キー, 項目, 記号, 長短, 文例)
# ---------------------------------------------------------------------------
BEHAVIOR_TPL = {
    "seikatsu": {
        ("◎", "long"): [
            "時間を守って行動し、身の回りの整理整頓も欠かさず、規則正しい生活を送{っていました|っていた}。",
            "いつも落ち着いて生活し、持ち物の準備や後片付けを自分から進んで行{っていました|っていた}。",
            "あいさつや返事が気持ちよくでき、けじめのある生活態度が身に付{いていました|いていた}。",
        ],
        ("○", "long"): [
            "時間を意識して行動し、身の回りを整えようと努め{ていました|ていた}。",
            "決められたことを守って生活しようとする姿が見られ{ました|た}。",
        ],
        ("◎", "short"): [
            "時間を守り、整理整頓を欠かさず生活し{ていました|ていた}。",
            "けじめのある生活態度が身に付{いていました|いていた}。",
        ],
        ("○", "short"): ["時間を意識し、生活を整えようと努め{ていました|ていた}。"],
    },
    "kenko": {
        ("◎", "long"): [
            "休み時間には友達と外で元気に体を動かし、健康な生活を心がけ{ていました|ていた}。",
            "運動に進んで参加し、繰り返し練習することで体力を高め{ていました|ていた}。",
            "手洗いや換気などを進んで行い、自分の健康を大切にし{ていました|ていた}。",
        ],
        ("○", "long"): [
            "外遊びや運動に進んで参加し、体を動かす楽しさを味わ{っていました|っていた}。",
            "健康に気を付けて生活しようとする姿が見られ{ました|た}。",
        ],
        ("◎", "short"): ["進んで運動に取り組み、健康な生活を心がけ{ていました|ていた}。"],
        ("○", "short"): ["体を動かす活動に進んで参加し{ていました|ていた}。"],
    },
    "jishu": {
        ("◎", "long"): [
            "自分がすべきことを考えて判断し、進んで行動することができ{ました|た}。",
            "よくないと思ったことは自分で考えて我慢し、正しく行動しようとする強さがあ{りました|った}。",
            "言われる前に気付いて動くことができ、周りのよい手本となって{いました|いた}。",
        ],
        ("○", "long"): [
            "自分で考えて行動しようとする姿が見られ{ました|た}。",
            "めあてをもって物事に取り組もうと努め{ていました|ていた}。",
        ],
        ("◎", "short"): ["自分で判断し、進んで行動することができ{ました|た}。"],
        ("○", "short"): ["自分で考えて行動しようと努め{ていました|ていた}。"],
    },
    "sekinin": {
        ("◎", "long"): [
            "任された仕事は、最後まで責任をもってやり遂げ{ていました|ていた}。",
            "自分の役割を自覚し、面倒なことも投げ出さずに最後までやり抜{いていました|いていた}。",
            "一度引き受けたことは必ずやり通し、周りからの信頼を集め{ていました|ていた}。",
        ],
        ("○", "long"): [
            "自分の仕事を最後まで行おうと努め{ていました|ていた}。",
            "任された役割に責任をもって取り組む姿が見られ{ました|た}。",
        ],
        ("◎", "short"): ["任された仕事を最後までやり遂げ{ていました|ていた}。"],
        ("○", "short"): ["自分の役割に責任をもって取り組{んでいました|んでいた}。"],
    },
    "soui": {
        ("◎", "long"): [
            "よりよい方法はないかと考え、工夫しながら課題に取り組{んでいました|んでいた}。",
            "自分なりのアイデアを出し、試行錯誤を重ねながら工夫を深め{ていました|ていた}。",
            "これまでの学びを生かして新しい方法を考え出し、実践する発想の豊かさがあ{りました|った}。",
        ],
        ("○", "long"): [
            "工夫して取り組もうとする姿が見られ{ました|た}。",
            "よりよくしようと考えながら活動し{ていました|ていた}。",
        ],
        ("◎", "short"): ["工夫を重ねて課題に取り組{んでいました|んでいた}。"],
        ("○", "short"): ["よりよい方法を考えて活動し{ていました|ていた}。"],
    },
    "omoiyari": {
        ("◎", "long"): [
            "困っている友達にそっと寄り添い、優しく声をかけ{ていました|ていた}。",
            "相手の気持ちを考えて行動し、友達と力を合わせて活動を進め{ていました|ていた}。",
            "だれにでも温かく接し、学級を和やかにする存在{でした|であった}。",
        ],
        ("○", "long"): [
            "友達と協力して活動しようとする姿が見られ{ました|た}。",
            "友達の気持ちを考えて接しようと努め{ていました|ていた}。",
        ],
        ("◎", "short"): ["友達に優しく接し、協力して活動し{ていました|ていた}。"],
        ("○", "short"): ["相手の気持ちを考えて行動しようと努め{ていました|ていた}。"],
    },
    "seimei": {
        ("◎", "long"): [
            "生き物や植物の世話を毎日続けることで、命の大切さを実感し{ていました|ていた}。",
            "自然に親しみ、身の回りの環境を大切にする行動を積み重ね{ていました|ていた}。",
        ],
        ("○", "long"): [
            "生き物や植物を大切に世話し{ていました|ていた}。",
            "自然や生命を大切にしようとする姿が見られ{ました|た}。",
        ],
        ("◎", "short"): ["生き物の世話を続け、命を大切にし{ていました|ていた}。"],
        ("○", "short"): ["自然や生命を大切にしようとし{ていました|ていた}。"],
    },
    "kinro": {
        ("◎", "long"): [
            "清掃活動に黙々と取り組み、みんなのために働く姿が立派{でした|であった}。",
            "人が嫌がる仕事も進んで引き受け、学級のために力を尽くし{ていました|ていた}。",
            "自分から仕事を見付けて働くことができ、その働きぶりは学級の手本となって{いました|いた}。",
        ],
        ("○", "long"): [
            "当番や清掃の仕事に真面目に取り組{んでいました|んでいた}。",
            "みんなのために働こうとする姿が見られ{ました|た}。",
        ],
        ("◎", "short"): ["みんなのために進んで働{いていました|いていた}。"],
        ("○", "short"): ["当番や清掃に真面目に取り組{んでいました|んでいた}。"],
    },
    "kosei": {
        ("◎", "long"): [
            "だれに対しても分け隔てなく接し、公平な態度で行動し{ていました|ていた}。",
            "相手の立場に立って考え、正しいと思うことを大切にして判断し{ていました|ていた}。",
        ],
        ("○", "long"): [
            "友達に対して分け隔てなく接しようと努め{ていました|ていた}。",
            "公平に判断しようとする姿が見られ{ました|た}。",
        ],
        ("◎", "short"): ["だれにでも分け隔てなく接し{ていました|ていた}。"],
        ("○", "short"): ["公平に判断しようと努め{ていました|ていた}。"],
    },
    "kokyo": {
        ("◎", "long"): [
            "学校のきまりを守り、みんなが使う物や場所を大切にし{ていました|ていた}。",
            "学級や学校のためになることを考えて行動し{ていました|ていた}。",
        ],
        ("○", "long"): [
            "きまりを守って生活しようとする姿が見られ{ました|た}。",
            "みんなで使う物を大切にしようと努め{ていました|ていた}。",
        ],
        ("◎", "short"): ["きまりを守り、公共の物を大切にし{ていました|ていた}。"],
        ("○", "short"): ["きまりを守って生活しようと努め{ていました|ていた}。"],
    },
}

SIMPLE_TPL = {
    "committee": {
        "long": [
            "{{委員会}}では、{{役割句}}{{活動句}}責任をもって取り組{んでいました|んでいた}。",
            "{{委員会}}では、{{役割句}}{{活動句}}こつこつと取り組み、学校のために働{いていました|いていた}。",
            "{{委員会}}では、{{役割句}}{{活動句}}進んで励{んでいました|んでいた}。",
            "{{委員会}}では、{{役割句}}{{活動句}}熱心に活動{していました|していた}。",
            "{{委員会}}の仕事では、{{役割句}}{{活動句}}最後まで責任を果たそうと力を尽くし{ていました|ていた}。",
        ],
        "short": [
            "{{委員会}}で{{役割句}}{{活動句}}責任をもって取り組{んでいました|んでいた}。",
            "{{委員会}}の活動に進んで参加{していました|していた}。",
            "{{委員会}}の仕事を最後まで頑張{っていました|っていた}。",
        ],
    },
    "kakari": {
        "long": [
            "係活動では、{{係}}として{{仕事句}}毎日欠かさず行{っていました|っていた}。",
            "{{係}}として、{{仕事句}}忘れずに行い、学級の生活を支え{ていました|ていた}。",
            "{{係}}の仕事に工夫を加えながら取り組み、みんなが気持ちよく過ごせるようにし{ていました|ていた}。",
            "{{係}}として、{{仕事句}}自分から進んで行う姿が見られ{ました|た}。",
            "{{係}}の仕事を最後まで丁寧にやり遂げ{ていました|ていた}。",
        ],
        "short": [
            "{{係}}として{{仕事句}}最後まで行{っていました|っていた}。",
            "{{係}}の仕事に責任をもって取り組{んでいました|んでいた}。",
            "{{係}}の仕事をいつも忘れずに続け{ていました|ていた}。",
        ],
    },
    "event": {
        "long": [
            "{{行事}}では、{{役割句}}{{頑張り句}}力を尽くし{ていました|ていた}。",
            "{{行事}}では、{{役割句}}{{頑張り句}}友達と協力しながら最後までやり遂げ{ました|た}。",
            "{{行事}}に向けて、{{役割句}}{{頑張り句}}練習や準備に粘り強く取り組{んでいました|んでいた}。",
            "{{行事}}では、{{役割句}}{{頑張り句}}精一杯の姿を見せてくれ{ました|た}。",
            "{{行事}}に向けて、{{役割句}}{{頑張り句}}本番まで一生懸命練習を重ね{ていました|ていた}。",
        ],
        "short": [
            "{{行事}}では{{役割句}}力を尽くし{ました|た}。",
            "{{行事}}に向けて粘り強く取り組{んでいました|んでいた}。",
            "{{行事}}で精一杯の姿を見せてくれ{ました|た}。",
        ],
    },
    "moral-intro": {
        "long": [
            "道徳の学習では、{{学期}}を通して自分の考えをもち、友達の意見を聞きながら新しい見方を広げ{ていました|ていた}。",
            "道徳の学習では、教材の登場人物の気持ちを自分に置き換えて考え{ていました|ていた}。",
        ],
        "short": ["道徳の学習では、自分の考えをもって話し合いに参加し{ていました|ていた}。"],
    },
    "moral-material": {
        "long": [
            "「{{教材}}」の学習では、{{項目句}}について自分の経験と重ねて考え、{{気付き句}}と振り返{っていました|っていた}。",
            "「{{教材}}」の学習では、友達の意見を聞いて違う見方を理解し、{{気付き句}}と自分の考えをまとめ{ていました|ていた}。",
            "「{{教材}}」では、{{項目句}}について様々な立場や見方から考え、{{気付き句}}と気づくことができ{ました|た}。",
        ],
        "short": [
            "「{{教材}}」の学習で、{{気付き句}}と振り返{っていました|っていた}。",
            "「{{教材}}」でいろいろな角度から{{項目句}}について考え{ていました|ていた}。",
            "「{{教材}}」を通して、{{気付き句}}と気づくことができ{ました|た}。",
        ],
    },
    "moral-material-nonote": {
        "long": [
            "「{{教材}}」の学習では、{{項目句}}について自分の経験と重ねて{考えていました|考えていた}。",
            "「{{教材}}」の学習では、友達の意見を聞きながら、{{項目句}}について違う考え方を知り{ました|た}。",
            "「{{教材}}」では、{{項目句}}について様々な立場から{考えていました|考えていた}。",
        ],
        "short": [
            "「{{教材}}」でいろいろな見方から{{項目句}}について考え{ていました|ていた}。",
            "「{{教材}}」を通して、{{項目句}}について自分の思いを見つめ直{していました|していた}。",
            "「{{教材}}」の学習で、{{項目句}}について改めて考え{ました|た}。",
        ],
    },
    "moral-closing": {
        "long": [
            "これからの生活に学んだことを生かそうとする意欲が高まって{います|いる}。",
            "学んだことを自分の行動に生かそうとする姿が見られ{ます|る}。",
        ],
        "short": ["学びを日常で実践しようとして{います|いる}。"],
    },
    "sougou": {
        "long": [
            "総合的な学習の時間の「{{単元}}」では、{{課題句}}自分で課題を設定し、{{活動句}}調べたことを整理して{まとめていました|まとめていた}。",
            "「{{単元}}」の学習では、{{活動句}}友達と協力して調べ、{{成果句}}分かりやすく伝えることができ{ました|た}。",
            "「{{単元}}」の探究では、{{課題句}}繰り返し調べ直し、{{成果句}}自分の発見をまとめることができ{ました|た}。",
        ],
        "short": [
            "「{{単元}}」で{{課題句}}調べ、{{成果句}}考えをまとめ{ていました|ていた}。",
            "「{{単元}}」で{{課題句}}進んで調べ、{{成果句}}成果を表現し{ていました|ていた}。",
        ],
    },
    "sougou-closing": {
        "long": [
            "学んだことを日々の生活に生かそうとする意欲が高まって{います|いる}。",
            "追究したことを次の学習につなげようとして{います|いる}。",
        ],
        "short": ["学びを生活に生かそうとして{います|いる}。"],
    },
    "closing": {
        "long": [
            "これからの活躍が楽しみ{です|である}。",
            "次の学期の頑張りにも期待して{います|いる}。",
            "この良さをさらに伸ばしていってほしいと{思います|思う}。",
        ],
    },
}

SIMPLE_LABEL = {
    "committee": "委員会",
    "kakari": "係",
    "event": "行事・行事の係",
    "moral-intro": "道徳（導入）",
    "moral-material": "道徳（教材・振り返りあり）",
    "moral-material-nonote": "道徳（教材・振り返りなし）",
    "moral-closing": "道徳（結び）",
    "sougou": "総合的な学習",
    "sougou-closing": "総合（結び）",
    "closing": "結びの一文",
}


# ---------------------------------------------------------------------------
def style_header(ws, row, ncols, width_map=None, height=22):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = HEAD_FONT
        cell.fill = HEAD_FILL
        cell.border = BORDER
        cell.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
    ws.row_dimensions[row].height = height
    if width_map:
        for col, w in width_map.items():
            ws.column_dimensions[col].width = w


def put_headers(ws, headers, width_map=None):
    for i, h in enumerate(headers, start=1):
        ws.cell(row=1, column=i, value=h)
    style_header(ws, 1, len(headers), width_map)
    ws.freeze_panes = "B2"


def add_list_validation(ws, rng, items):
    dv = DataValidation(type="list", formula1='"%s"' % ",".join(items), allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(rng)
    return dv


# ---------------------------------------------------------------------------
def sheet_intro(wb):
    ws = wb.create_sheet("はじめに")
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 104

    lines = [
        ("title", "通信表・指導要録 所見作成支援（エクセル版）"),
        ("", ""),
        ("head", "■ 児童生徒の個人情報は入力しません"),
        ("", "管理は出席番号のみです。氏名・住所などの入力欄は設けていません。"),
        ("", "所見の中に氏名らしい表現（〜さん・くん）があると「所見」シートの警告欄に表示します。"),
        ("", "端末を手放すときや年度が変わったときは、マクロ「ClearAllData」で入力を消してください。"),
        ("", ""),
        ("head", "■ はじめに一度だけ：マクロの取り込み"),
        ("", "このファイルはマクロ無し（.xlsx）です。次の手順で一度だけマクロを入れてください。"),
        ("", "取り込んだあとに .xlsm で保存すれば、そのファイルはそのまま配布して使えます。"),
        ("", ""),
        ("step", "1. このブックを「Excel マクロ有効ブック (*.xlsm)」として保存し直す"),
        ("step", "2. Alt + F11 を押して VBE（Visual Basic Editor）を開く"),
        ("step", "3. メニューの「ファイル」→「ファイルのインポート」を選ぶ"),
        ("step", "4. vba フォルダの中の 4 つのファイルを順に取り込む"),
        ("", "     modUtil.bas / modTemplates.bas / modGenerate.bas / modApp.bas"),
        ("step", "5. VBE を閉じて、上書き保存する"),
        ("", ""),
        ("head", "■ 使い方"),
        ("step", "1.「設定」  学期・出席番号の最大・文字数・所見に入れる項目と順番を決める"),
        ("step", "2.「委員会・係」「行事」「行動の記録」「見取り」「道徳」「総合」に入力する"),
        ("", "     エクセルの他の表からコピーして貼り付けられます。"),
        ("step", "3. Alt + F8 →「GenerateAll」を実行する（ボタンを置いても構いません）"),
        ("step", "4.「所見」シートで確認・手直しする"),
        ("", ""),
        ("head", "■「所見」シートの使い方"),
        ("", "別案 … 数字を変えると言い回しが変わります（0,1,2…）。"),
        ("", "手直し … 「はい」にした行は、作り直しても上書きされません。"),
        ("", "警告 … 文字数超過・未使用の文・氏名らしい表現をお知らせします。"),
        ("", ""),
        ("head", "■ 文例を変えたいとき"),
        ("", "「文例」シートを直すだけで、生成される文が変わります（VBAの編集は不要です）。"),
        ("", "  {敬体|常体} … 通信表と指導要録での出し分け"),
        ("", "  {{キー}}     … 入力内容の差し込み"),
        ("", "同じ区分の文例を増やすほど、同じ言い回しが続きにくくなります。"),
        ("", ""),
        ("head", "■ 注意"),
        ("", "生成される文は文例の組み合わせです。そのまま提出せず、"),
        ("", "必ず担任がその子の事実と照らして確認・修正してください。"),
        ("", "「そのまま使う」の見取りを常体へ変換する処理は機械的です。要録に使う前に読み返してください。"),
    ]
    r = 1
    for kind, text in lines:
        cell = ws.cell(row=r, column=2, value=text)
        if kind == "title":
            cell.font = TITLE_FONT
            ws.row_dimensions[r].height = 26
        elif kind == "head":
            cell.font = Font(bold=True, size=11, color=ACCENT)
        elif kind == "step":
            cell.font = Font(size=11)
        else:
            cell.font = NOTE_FONT
        cell.alignment = Alignment(vertical="center")
        r += 1
    return ws


def sheet_settings(wb):
    ws = wb.create_sheet("設定")
    ws.cell(row=1, column=1, value="項目")
    ws.cell(row=1, column=2, value="キー")
    ws.cell(row=1, column=3, value="値")
    ws.cell(row=1, column=5, value="所見に入れる項目（上ほど優先）")
    ws.cell(row=1, column=6, value="キー")
    ws.cell(row=1, column=7, value="使う")
    ws.cell(row=1, column=9, value="行動の記録の優先順")
    ws.cell(row=1, column=10, value="キー")
    style_header(ws, 1, 10, {
        "A": 34, "B": 20, "C": 14, "D": 2,
        "E": 26, "F": 14, "G": 8, "H": 2, "I": 24, "J": 14,
    })
    ws.freeze_panes = "A2"

    rows = [
        ("学期", "term", "2学期"),
        ("出席番号の最大", "count", 35),
        ("通信表の文字数（行動・活動・見取り）", "limitReport", 120),
        ("指導要録の文字数（行動・活動・見取り）", "limitYouroku", 60),
        ("道徳 通信表の文字数", "moralLimitReport", 120),
        ("道徳 指導要録の文字数", "moralLimitYouroku", 60),
        ("総合 通信表の文字数", "sougouLimitReport", 120),
        ("総合 指導要録の文字数", "sougouLimitYouroku", 60),
        ("行動の記録から書く数", "behaviorMax", 2),
        ("行動の選び方 order=並び順 / mark=◎優先", "behaviorPick", "order"),
        ("結びの一文を入れる", "useClosing", "いいえ"),
        ("道徳の導入文を入れる", "moralIntro", "はい"),
        ("道徳の結びを入れる", "moralClosing", "はい"),
        ("総合の結びを入れる", "sougouClosing", "はい"),
    ]
    for i, (label, key, val) in enumerate(rows, start=2):
        ws.cell(row=i, column=1, value=label).border = BORDER
        ws.cell(row=i, column=2, value=key).border = BORDER
        c = ws.cell(row=i, column=3, value=val)
        c.border = BORDER
        c.alignment = Alignment(horizontal="center")

    for i, (key, label) in enumerate(CATEGORIES, start=2):
        ws.cell(row=i, column=5, value=label).border = BORDER
        ws.cell(row=i, column=6, value=key).border = BORDER
        c = ws.cell(row=i, column=7, value="はい")
        c.border = BORDER
        c.alignment = Alignment(horizontal="center")

    for i, (key, label) in enumerate(BEHAVIOR_ITEMS, start=2):
        ws.cell(row=i, column=9, value=label).border = BORDER
        ws.cell(row=i, column=10, value=key).border = BORDER

    add_list_validation(ws, "C12:C15", ["はい", "いいえ"])
    add_list_validation(ws, "G2:G6", ["はい", "いいえ"])
    add_list_validation(ws, "C11", ["order", "mark"])

    ws.cell(row=18, column=1,
            value="※ 行の上下を入れ替えると優先順が変わります（文字数に収まらない分は下から削られます）。"
            ).font = NOTE_FONT
    ws.cell(row=19, column=1,
            value="※「行動の記録」シートの列の並びは固定です。ここでの順は「どれを先に書くか」に効きます。"
            ).font = NOTE_FONT
    return ws


def sheet_roster(wb, count):
    ws = wb.create_sheet("委員会・係")
    put_headers(ws, ["番号", "委員会", "委員会での役割", "委員会の活動内容", "係", "係の仕事"],
                {"A": 6, "B": 18, "C": 18, "D": 26, "E": 16, "F": 26})
    sample = [
        (1, "図書委員会", "副委員長", "本の貸し出し", "黒板係", "黒板をきれいにする仕事"),
        (2, "放送委員会", "", "昼の放送", "配り係", "プリントを配る仕事"),
        (3, "体育委員会", "委員長", "用具の準備", "生き物係", "メダカの水替え"),
        (4, "図書委員会", "", "本の整理", "電気係", "教室の電気を管理する仕事"),
    ]
    for no in range(1, count + 1):
        ws.cell(row=no + 1, column=1, value=no).alignment = Alignment(horizontal="center")
    for row in sample:
        r = row[0] + 1
        for i, v in enumerate(row[1:], start=2):
            ws.cell(row=r, column=i, value=v)
    return ws


def sheet_event(wb, count):
    ws = wb.create_sheet("行事")
    headers = ["番号"]
    for i in (1, 2, 3):
        headers += [f"行事{i}", f"行事{i}の係", f"行事{i}の頑張り"]
    widths = {"A": 6}
    for i, col in enumerate("BCDEFGHIJ"):
        widths[col] = [16, 16, 26][i % 3]
    put_headers(ws, headers, widths)
    for no in range(1, count + 1):
        ws.cell(row=no + 1, column=1, value=no).alignment = Alignment(horizontal="center")
    samples = {
        1: ("運動会", "応援団", "大きな声で応援し"),
        2: ("学習発表会", "司会", "台本を何度も読み込み"),
        3: ("運動会", "リレー選手", "バトンの練習を重ね"),
        4: ("遠足", "班長", "班のみんなに声をかけ"),
    }
    for no, (a, b, c) in samples.items():
        ws.cell(row=no + 1, column=2, value=a)
        ws.cell(row=no + 1, column=3, value=b)
        ws.cell(row=no + 1, column=4, value=c)
    return ws


def sheet_behavior(wb, count):
    ws = wb.create_sheet("行動の記録")
    headers = ["番号"] + [label for _, label in BEHAVIOR_ITEMS]
    widths = {"A": 6}
    for i in range(len(BEHAVIOR_ITEMS)):
        widths[get_column_letter(i + 2)] = 11
    put_headers(ws, headers, widths)
    for no in range(1, count + 1):
        ws.cell(row=no + 1, column=1, value=no).alignment = Alignment(horizontal="center")
        for c in range(2, len(BEHAVIOR_ITEMS) + 2):
            ws.cell(row=no + 1, column=c).alignment = Alignment(horizontal="center")
    last = get_column_letter(len(BEHAVIOR_ITEMS) + 1)
    add_list_validation(ws, f"B2:{last}{count + 1}", [MARK_A, MARK_B])
    samples = {
        1: {"sekinin": MARK_A, "omoiyari": MARK_B},
        2: {"seikatsu": MARK_A, "kinro": MARK_A},
        3: {"kenko": MARK_A, "jishu": MARK_B},
        4: {"seimei": MARK_A, "soui": MARK_B},
    }
    keyidx = {k: i + 2 for i, (k, _) in enumerate(BEHAVIOR_ITEMS)}
    for no, marks in samples.items():
        for k, mk in marks.items():
            ws.cell(row=no + 1, column=keyidx[k], value=mk).alignment = Alignment(horizontal="center")
    return ws


def sheet_obs(wb, count):
    ws = wb.create_sheet("見取り")
    headers = ["番号"]
    for i in (1, 2, 3):
        headers += [f"見取り{i}", f"文型{i}"]
    put_headers(ws, headers, {"A": 6, "B": 34, "C": 22, "D": 34, "E": 22, "F": 34, "G": 22})
    for no in range(1, count + 1):
        ws.cell(row=no + 1, column=1, value=no).alignment = Alignment(horizontal="center")
    labels = [lab for _, lab, _, _ in OBS_PATTERNS]
    for col in ("C", "E", "G"):
        add_list_validation(ws, f"{col}2:{col}{count + 1}", labels)
    samples = {
        1: ("友達の失敗を責めずに励ます", labels[0]),
        2: ("自分の考えを理由を付けて発表する", labels[1]),
        3: ("当番の仕事を人一倍ていねいに行っ", labels[2]),
        4: ("学級のムードメーカー", labels[3]),
    }
    for no, (tx, lab) in samples.items():
        ws.cell(row=no + 1, column=2, value=tx)
        ws.cell(row=no + 1, column=3, value=lab)
    ws.cell(row=count + 3, column=2,
            value="※ 文型に合う形で書きます（例:「〜ていました」なら「…行っ」で止める）。").font = NOTE_FONT
    return ws


def sheet_moral(wb):
    ws = wb.create_sheet("道徳")
    put_headers(ws, ["番号", "教材名", "内容項目", "児童の振り返り"],
                {"A": 6, "B": 24, "C": 30, "D": 46})
    dv = DataValidation(type="list", formula1="='文例'!$H$2:$H$23", allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("C2:C400")
    samples = [
        (1, "ぼくのボール", "親切、思いやり", "相手の気持ちを考えて行動したい"),
        (2, "お母さんのせいきゅう書", "家族愛、家庭生活の充実", "家族のためにできることをしたい"),
        (3, "とべないホタル", "生命の尊さ", "命はどれも大切だと思った"),
    ]
    for i, row in enumerate(samples, start=2):
        for j, v in enumerate(row, start=1):
            ws.cell(row=i, column=j, value=v)
    ws.cell(row=8, column=2,
            value="※ 1人につき何行でも書けます。振り返りが空欄のときは、気付きを作らない文例を使います。"
            ).font = NOTE_FONT
    return ws


def sheet_sougou(wb):
    ws = wb.create_sheet("総合")
    put_headers(ws, ["番号", "単元名", "探究課題", "活動の様子", "成果"],
                {"A": 6, "B": 24, "C": 26, "D": 30, "E": 30})
    samples = [
        (1, "米づくりを調べよう", "地域の農業", "農家の方に話を聞き", "分かったことを新聞にまとめ"),
        (2, "町のバリアフリー", "だれもが暮らしやすい町", "町を歩いて調べ", "提案をポスターにして"),
    ]
    for i, row in enumerate(samples, start=2):
        for j, v in enumerate(row, start=1):
            ws.cell(row=i, column=j, value=v)
    ws.cell(row=7, column=2, value="※ 1人につき何行でも書けます。").font = NOTE_FONT
    return ws


def sheet_templates(wb):
    ws = wb.create_sheet("文例")
    put_headers(ws, ["区分（説明）", "キー", "項目", "記号", "長短", "文例"],
                {"A": 30, "B": 22, "C": 14, "D": 6, "E": 8, "F": 96})
    ws.column_dimensions["H"].width = 34

    r = 2
    for key, label in BEHAVIOR_ITEMS:
        for (mk, variant), texts in BEHAVIOR_TPL[key].items():
            for t in texts:
                ws.cell(row=r, column=1, value=f"行動の記録 / {label}")
                ws.cell(row=r, column=2, value="behavior")
                ws.cell(row=r, column=3, value=key)
                ws.cell(row=r, column=4, value=mk)
                ws.cell(row=r, column=5, value=variant)
                ws.cell(row=r, column=6, value=t)
                r += 1

    for key in ["committee", "kakari", "event", "moral-intro", "moral-material",
                "moral-material-nonote", "moral-closing", "sougou", "sougou-closing", "closing"]:
        for variant, texts in SIMPLE_TPL[key].items():
            for t in texts:
                ws.cell(row=r, column=1, value=SIMPLE_LABEL[key])
                ws.cell(row=r, column=2, value=key)
                ws.cell(row=r, column=5, value=variant)
                ws.cell(row=r, column=6, value=t)
                r += 1

    for key, label, tpl, hint in OBS_PATTERNS:
        ws.cell(row=r, column=1, value=label)
        ws.cell(row=r, column=2, value="obs")
        ws.cell(row=r, column=3, value=key)
        ws.cell(row=r, column=6, value=tpl)
        r += 1

    for row in range(2, r):
        ws.cell(row=row, column=6).alignment = Alignment(wrap_text=True, vertical="center")

    ws.cell(row=1, column=8, value="道徳の内容項目").font = HEAD_FONT
    ws.cell(row=1, column=8).fill = HEAD_FILL
    for i, v in enumerate(MORAL_VALUES, start=2):
        ws.cell(row=i, column=8, value=v)

    ws.freeze_panes = "C2"
    return ws


def sheet_output(wb, count):
    ws = wb.create_sheet("所見")
    headers = ["番号", "別案",
               "通信表 所見（敬体）", "字数",
               "指導要録 所見（常体）", "字数",
               "道徳（通信表）", "字数",
               "道徳（指導要録）", "字数",
               "総合（通信表）", "字数",
               "総合（指導要録）", "字数",
               "警告", "手直し"]
    widths = {"A": 6, "B": 6, "C": 46, "D": 6, "E": 34, "F": 6, "G": 40, "H": 6,
              "I": 32, "J": 6, "K": 40, "L": 6, "M": 32, "N": 6, "O": 34, "P": 8}
    put_headers(ws, headers, widths)
    for no in range(1, count + 1):
        r = no + 1
        ws.cell(row=r, column=1, value=no).alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=2, value=0).alignment = Alignment(horizontal="center")
        for c in (3, 5, 7, 9, 11, 13, 15):
            ws.cell(row=r, column=c).alignment = Alignment(wrap_text=True, vertical="top")
    add_list_validation(ws, f"P2:P{count + 1}", ["はい", "いいえ"])
    return ws


def copy_vba():
    """VBA モジュールを Shift-JIS(cp932)・CRLF で dist/vba に書き出す。
    VBE の「ファイルのインポート」は OS の既定コードページで読むため、
    日本語 Windows でコメントが化けないようにそろえておく。"""
    import glob
    src_dir = os.path.join(os.path.dirname(OUT_DIR), "vba")
    dst_dir = os.path.join(OUT_DIR, "vba")
    os.makedirs(dst_dir, exist_ok=True)
    for path in sorted(glob.glob(os.path.join(src_dir, "*.bas"))):
        text = open(path, encoding="utf-8").read().replace("\n", "\r\n")
        dst = os.path.join(dst_dir, os.path.basename(path))
        with open(dst, "wb") as f:
            f.write(text.encode("cp932"))
        print("  vba:", os.path.basename(dst))


def main():
    count = 35
    wb = Workbook()
    wb.remove(wb.active)

    sheet_intro(wb)
    sheet_settings(wb)
    sheet_roster(wb, count)
    sheet_event(wb, count)
    sheet_behavior(wb, count)
    sheet_obs(wb, count)
    sheet_moral(wb)
    sheet_sougou(wb)
    sheet_templates(wb)
    sheet_output(wb, count)

    os.makedirs(OUT_DIR, exist_ok=True)
    wb.save(OUT_PATH)
    print("wrote:", OUT_PATH)
    for ws in wb.worksheets:
        print(f"  {ws.title}: {ws.max_row} rows x {ws.max_column} cols")
    copy_vba()


if __name__ == "__main__":
    main()
