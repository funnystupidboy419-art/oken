#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
VBA へ移植した生成ロジックの検証。

  VBA と同じ手順を Python で書き起こし、元の JavaScript 版（node 実行）の
  出力と一文字ずつ突き合わせる。文例は生成した .xlsx の「文例」シートから
  読むので、シート -> ロジック の経路もまとめて確かめられる。

  python3 tools/verify_logic.py
"""
import json
import os
import re
import subprocess
import sys

from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(ROOT, "dist", "所見作成支援.xlsx")

MARK_A = "◎"
MARK_B = "○"
BEHAVIOR_ORDER = ["seikatsu", "kenko", "jishu", "sekinin", "soui",
                  "omoiyari", "seimei", "kinro", "kosei", "kokyo"]
CATEGORY_ORDER = ["observation", "event", "committee", "kakari", "behavior"]

SETTINGS = {
    "term": "2学期",
    "order": CATEGORY_ORDER,
    "enabled": {k: True for k in CATEGORY_ORDER},
    "behaviorOrder": BEHAVIOR_ORDER,
    "behaviorMax": 2,
    "behaviorPick": "order",
    "limitReport": 120, "limitYouroku": 60,
    "moralLimitReport": 120, "moralLimitYouroku": 60,
    "sougouLimitReport": 120, "sougouLimitYouroku": 60,
    "useClosing": False,
    "moralIntro": True, "moralClosing": True, "sougouClosing": True,
}

STUDENTS = [
    {"no": 1,
     "committee": "図書委員会", "committeeRole": "副委員長", "committeeWork": "本の貸し出し",
     "kakari": "黒板係", "kakariWork": "黒板をきれいにする仕事",
     "events": [{"name": "運動会", "role": "応援団", "effort": "大きな声で応援し"}],
     "behavior": {"sekinin": MARK_A, "omoiyari": MARK_B},
     "observations": [{"text": "友達の失敗を責めずに励ます", "pattern": "sugata"}],
     "moral": [{"material": "ぼくのボール", "value": "親切、思いやり",
                "note": "相手の気持ちを考えて行動したい"}],
     "sougou": [{"unit": "米づくりを調べよう", "theme": "地域の農業",
                 "activity": "農家の方に話を聞き", "result": "分かったことを新聞にまとめ"}]},
    {"no": 2,
     "committee": "放送委員会", "committeeRole": "", "committeeWork": "昼の放送",
     "kakari": "配り係", "kakariWork": "プリントを配る仕事",
     "events": [{"name": "学習発表会", "role": "司会", "effort": "台本を何度も読み込み"}],
     "behavior": {"seikatsu": MARK_A, "kinro": MARK_A},
     "observations": [{"text": "自分の考えを理由を付けて発表する", "pattern": "dekita"}],
     "moral": [{"material": "お母さんのせいきゅう書", "value": "家族愛、家庭生活の充実",
                "note": "家族のためにできることをしたい"}],
     "sougou": [{"unit": "町のバリアフリー", "theme": "だれもが暮らしやすい町",
                 "activity": "町を歩いて調べ", "result": "提案をポスターにして"}]},
    {"no": 3,
     "committee": "体育委員会", "committeeRole": "委員長", "committeeWork": "用具の準備",
     "kakari": "生き物係", "kakariWork": "メダカの水替え",
     "events": [{"name": "運動会", "role": "リレー選手", "effort": "バトンの練習を重ね"}],
     "behavior": {"kenko": MARK_A, "jishu": MARK_B},
     "observations": [{"text": "当番の仕事を人一倍ていねいに行っ", "pattern": "teita"}],
     "moral": [{"material": "とべないホタル", "value": "生命の尊さ",
                "note": "命はどれも大切だと思った"}],
     "sougou": []},
    {"no": 4,
     "committee": "図書委員会", "committeeRole": "", "committeeWork": "本の整理",
     "kakari": "電気係", "kakariWork": "教室の電気を管理する仕事",
     "events": [{"name": "遠足", "role": "班長", "effort": "班のみんなに声をかけ"}],
     "behavior": {"seimei": MARK_A, "soui": MARK_B},
     "observations": [{"text": "学級のムードメーカー", "pattern": "datta"}],
     "moral": [{"material": "とべないホタル", "value": "生命の尊さ", "note": ""}],
     "sougou": []},
]

# ---------------------------------------------------------------------------
# VBA と同じ手順（modUtil / modGenerate の書き起こし）
# ---------------------------------------------------------------------------
RE_STYLE = re.compile(r"\{([^{}|]*)\|([^{}|]*)\}")
RE_VAR = re.compile(r"\{\{([^{}]+)\}\}")


def tidy(s):
    s = s.replace(" ", "").replace("　", "")
    while "、、" in s:
        s = s.replace("、、", "、")
    while "。。" in s:
        s = s.replace("。。", "。")
    s = s.replace("、。", "。")
    s = s.lstrip("、。")
    return s.strip()


def render_tpl(tpl, vars_, style):
    if not tpl:
        return ""
    s = RE_STYLE.sub(lambda m: m.group(2) if style == "plain" else m.group(1), tpl)
    s = RE_VAR.sub(lambda m: str(vars_.get(re.sub(r"\s", "", m.group(1)), "")), s)
    return tidy(s)


def with_particle(text, particle):
    t = (text or "").strip()
    if not t:
        return ""
    if t.endswith(particle):
        return t
    if len(particle) == 1 and t[-1] in "をにでへとやのはがも、。":
        return t
    return t + particle


def suffix_str(text, tail):
    t = (text or "").strip()
    return t + tail if t else ""


def ending_key(text):
    return text[-9:] if len(text) > 9 else text


def render_candidate(part, tpl, style):
    t = render_tpl(tpl, part["vars"], style)
    if part.get("raw") and style == "plain":
        raise AssertionError("この検証データでは toPlain を通らないはず")
    if t and t[-1] not in "。」）)":
        t += "。"
    return t


def render_part_varied(part, use_short, style, bump, prev_ending):
    lst = part["short"] if use_short else part["long"]
    if not lst:
        return "", prev_ending
    n = len(lst)
    start = ((part["seed"] + bump) % n + n) % n
    fallback = None
    for k in range(n):
        t = render_candidate(part, lst[(start + k) % n], style)
        if fallback is None:
            fallback = t
        if t and (not prev_ending or ending_key(t) != prev_ending):
            return t, ending_key(t)
    return fallback, ending_key(fallback)


def assemble(parts, limit, style, bump):
    rendered = []
    prev_long = prev_short = ""
    for p in parts:
        lt, le = render_part_varied(p, False, style, bump, prev_long)
        st, se = render_part_varied(p, True, style, bump, prev_short)
        if lt:
            prev_long = le
        if st:
            prev_short = se
        if lt or st:
            rendered.append({"longText": lt, "shortText": st, "useShort": False})

    def total(upto):
        return sum(len((it["shortText"] if it["useShort"] else it["longText"]).replace("\n", ""))
                   for it in rendered[:upto])

    i = len(rendered)
    while i >= 1:
        if limit <= 0 or total(len(rendered)) <= limit:
            break
        it = rendered[i - 1]
        if it["shortText"] and it["shortText"] != it["longText"]:
            it["useShort"] = True
        i -= 1

    kept = len(rendered)
    while limit > 0 and kept > 1 and total(kept) > limit:
        kept -= 1

    texts = [(it["shortText"] if it["useShort"] else it["longText"]) for it in rendered[:kept]]
    return "".join(texts)


# ---------------------------------------------------------------------------
def load_templates():
    ws = load_workbook(XLSX)["文例"]
    d = {}
    for r in range(2, ws.max_row + 1):
        key = ws.cell(r, 2).value
        if not key:
            continue
        dk = (key, ws.cell(r, 3).value or "", ws.cell(r, 4).value or "", ws.cell(r, 5).value or "")
        d.setdefault(dk, []).append(ws.cell(r, 6).value)
    return d


def tpl_pair(t, key, item="", mk="", want_short=False):
    a = t.get((key, item, mk, "short" if want_short else "long"), [])
    b = t.get((key, item, mk, "long" if want_short else "short"), [])
    return a if a else b


def behavior_tpl(t, key, mk, want_short):
    col = tpl_pair(t, "behavior", key, mk, want_short)
    return col if col else tpl_pair(t, "behavior", key, MARK_B, want_short)


def part(cat, longs, shorts, vars_, seed, raw=False):
    return {"cat": cat, "long": longs, "short": shorts, "vars": vars_, "seed": seed, "raw": raw}


def build_main(st, t, mode, bump=0):
    style = "plain" if mode == "youroku" else "polite"
    limit = SETTINGS["limitYouroku"] if mode == "youroku" else SETTINGS["limitReport"]
    seed = st["no"]
    parts = []
    for k in SETTINGS["order"]:
        if not SETTINGS["enabled"].get(k, True):
            continue
        if k == "committee" and st["committee"]:
            v = {"委員会": st["committee"],
                 "役割句": suffix_str(st["committeeRole"], "として、"),
                 "活動句": with_particle(st["committeeWork"], "に") if st["committeeWork"] else "活動に"}
            parts.append(part("committee", tpl_pair(t, "committee"), tpl_pair(t, "committee", want_short=True), v, seed))
        elif k == "kakari" and st["kakari"]:
            v = {"係": st["kakari"],
                 "仕事句": with_particle(st["kakariWork"], "を") if st["kakariWork"] else "自分の仕事を"}
            parts.append(part("kakari", tpl_pair(t, "kakari"), tpl_pair(t, "kakari", want_short=True), v, seed))
        elif k == "event":
            for i, ev in enumerate(st["events"]):
                if not ev["name"]:
                    continue
                v = {"行事": ev["name"], "役割句": suffix_str(ev["role"], "として、"),
                     "頑張り句": suffix_str(ev["effort"], "、")}
                parts.append(part("event", tpl_pair(t, "event"), tpl_pair(t, "event", want_short=True), v, seed + i))
        elif k == "behavior":
            picked = [x for x in SETTINGS["behaviorOrder"] if st["behavior"].get(x) in (MARK_A, MARK_B)]
            if SETTINGS["behaviorPick"] == "mark":
                picked = ([x for x in picked if st["behavior"][x] == MARK_A] +
                          [x for x in picked if st["behavior"][x] == MARK_B])
            for i, bk in enumerate(picked[:SETTINGS["behaviorMax"]]):
                mk = st["behavior"][bk]
                parts.append(part("behavior", behavior_tpl(t, bk, mk, False),
                                  behavior_tpl(t, bk, mk, True), {}, seed + i))
        elif k == "observation":
            for i, ob in enumerate(st["observations"]):
                if not ob["text"]:
                    continue
                tpl = t.get(("obs", ob["pattern"], "", ""), t.get(("obs", "sugata", "", ""), [""]))[0]
                parts.append(part("observation", [tpl], [tpl], {"見取り": ob["text"]},
                                  seed + i, ob["pattern"] == "as-is"))
    return assemble(parts, limit, style, bump)


def build_moral(st, t, mode, bump=0):
    style = "plain" if mode == "youroku" else "polite"
    limit = SETTINGS["moralLimitYouroku"] if mode == "youroku" else SETTINGS["moralLimitReport"]
    seed = st["no"]
    parts = []
    if SETTINGS["moralIntro"]:
        parts.append(part("moral-intro", tpl_pair(t, "moral-intro"),
                          tpl_pair(t, "moral-intro", want_short=True),
                          {"学期": SETTINGS["term"]}, seed + bump))
    for i, m in enumerate(st["moral"]):
        if not m["material"]:
            continue
        note = (m["note"] or "").strip()
        key = "moral-material" if note else "moral-material-nonote"
        v = {"教材": m["material"], "項目句": m["value"].strip() or "大切なこと",
             "気付き句": "「" + note.rstrip("。") + "」" if note else ""}
        parts.append(part("moral", tpl_pair(t, key), tpl_pair(t, key, want_short=True), v, seed + i + bump))
    if SETTINGS["moralClosing"] and parts:
        parts.append(part("moral-closing", tpl_pair(t, "moral-closing"),
                          tpl_pair(t, "moral-closing", want_short=True), {}, seed + bump))
    return assemble(parts, limit, style, 0)


def build_sougou(st, t, mode, bump=0):
    style = "plain" if mode == "youroku" else "polite"
    limit = SETTINGS["sougouLimitYouroku"] if mode == "youroku" else SETTINGS["sougouLimitReport"]
    seed = st["no"]
    parts = []
    for i, g in enumerate(st["sougou"]):
        if not g["unit"]:
            continue
        v = {"単元": g["unit"],
             "課題句": with_particle(g["theme"], "について") if g["theme"] else "",
             "活動句": suffix_str(g["activity"], "、"), "成果句": suffix_str(g["result"], "、")}
        parts.append(part("sougou", tpl_pair(t, "sougou"), tpl_pair(t, "sougou", want_short=True),
                          v, seed + i + bump))
    if SETTINGS["sougouClosing"] and parts:
        parts.append(part("sougou-closing", tpl_pair(t, "sougou-closing"),
                          tpl_pair(t, "sougou-closing", want_short=True), {}, seed + bump))
    return assemble(parts, limit, style, 0)


# ---------------------------------------------------------------------------
JS_DRIVER = r"""
global.window = {};
require('%s/js/util.js');
require('%s/js/templates.js');
require('%s/js/generate.js');
const G = global.window.Generator;
const input = JSON.parse(process.argv[2]);
const out = [];
for (const st of input.students) {
  out.push({
    no: st.no,
    report:    G.build(st, input.settings, 'report').text,
    youroku:   G.build(st, input.settings, 'youroku').text,
    moralR:    G.buildMoral(st, input.settings, 'report').text,
    moralY:    G.buildMoral(st, input.settings, 'youroku').text,
    sougouR:   G.buildSougou(st, input.settings, 'report').text,
    sougouY:   G.buildSougou(st, input.settings, 'youroku').text,
  });
}
console.log(JSON.stringify(out));
""" % (ROOT, ROOT, ROOT)


def run_js():
    drv = "/tmp/js_gen_driver.js"
    with open(drv, "w") as f:
        f.write(JS_DRIVER)
    payload = json.dumps({"settings": SETTINGS, "students": STUDENTS}, ensure_ascii=False)
    r = subprocess.run(["node", drv, payload], capture_output=True, text=True)
    if r.returncode != 0:
        print(r.stderr)
        sys.exit(1)
    return json.loads(r.stdout)


def main():
    t = load_templates()
    js = {x["no"]: x for x in run_js()}

    fields = [("report", build_main, "report"), ("youroku", build_main, "youroku"),
              ("moralR", build_moral, "report"), ("moralY", build_moral, "youroku"),
              ("sougouR", build_sougou, "report"), ("sougouY", build_sougou, "youroku")]

    bad = 0
    checked = 0
    for st in STUDENTS:
        for name, fn, mode in fields:
            mine = fn(st, t, mode)
            theirs = js[st["no"]][name]
            checked += 1
            if mine != theirs:
                bad += 1
                print(f"--- 不一致  番号{st['no']} / {name}")
                print(f"  JS : {theirs}")
                print(f"  VBA: {mine}")
    print(f"\n照合 {checked} 件 / 不一致 {bad} 件")
    if bad == 0:
        print("生成ロジックは JavaScript 版と一致しました。")
        for st in STUDENTS[:2]:
            print(f"\n[番号{st['no']}] 通信表: {build_main(st, t, 'report')}")
            print(f"          指導要録: {build_main(st, t, 'youroku')}")
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main())
